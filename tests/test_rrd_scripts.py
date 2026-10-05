"""Regression tests for the RRD scripts (v2). Run: python -m unittest discover -s tests

R01-R13 turn the first external review's probes (P01-P13) into tests that
expect the corrected behaviour. Synthetic text only; not research material.
"""

from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
import trace_scan  # noqa: E402

ZH_BASE = """# 引言

城市更新研究通常认为，拆迁补偿标准决定了居民的去留[1]。本文认为，这一判断忽略了居民在补偿谈判之外的安排。

# 3 案例分析

## 3.1 补偿谈判

“我们家三代都住在这条巷子里，钱再多也买不回邻居。”（R3，2021）一位老住户这样解释她为何拒签。她最后在第四轮谈判后才签约，补偿款比首轮高出38%。这说明谈判轮次改变了补偿结果。

# 5 结论

本文发现，居民通过延长谈判改变了补偿结果。
"""

ZH_DEFENSIVE = """# 引言

城市更新研究通常认为，拆迁补偿标准决定了居民的去留[1]。本研究关注居民在补偿过程中的相关安排，而非仅关注补偿标准。

# 3 案例分析

## 3.1 补偿谈判

据受访者介绍整理，居民（R3）在第四轮谈判后签约。事件顺序依其明确叙述整理。居民的选择可能与家庭关系有关，这一调整与谈判安排、补偿条件和组织关系有关。由此，补偿安排与居民关系联系起来。

# 5 结论

居民去留与补偿安排及谈判条件有关，而不能仅由补偿标准推知。
"""

EN_TEXT = """# Introduction

Studies of platform labour often assume that algorithmic control leaves workers with little room to act (Smith, 2019). We argue that couriers build their own dispatch rules.

# Findings

"I never take orders from the north side after six, the bridge eats twenty minutes," one courier explained. Couriers who shared such rules earned 14% more per shift. As recorded in the transcript, the interviewee then adds a qualification. This is not to say that control is absent.
"""


def write(tmp: Path, name: str, text: str) -> Path:
    p = tmp / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")
    return p


def run_gate(tmp: Path, base, rev, rows=None, memo: str | None = None, extra=None):
    b = base if isinstance(base, Path) else write(tmp, "base.md", base)
    r = rev if isinstance(rev, Path) else write(tmp, "rev.md", rev)
    args = [sys.executable, str(SCRIPTS / "defense_gate.py"), "--baseline", str(b), "--revised", str(r),
            "--config", str(ROOT / "assets/templates/gate_config.default.json"),
            "--out-json", str(tmp / "g.json"), "--out-md", str(tmp / "g.md")]
    if rows is not None:
        reg = tmp / "reg.csv"
        with reg.open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=["sentence_excerpt", "cue_type", "evidence_locator", "why_reader_needs_it", "section"])
            w.writeheader()
            w.writerows(rows)
        args += ["--register", str(reg)]
    if memo is not None:
        args += ["--route-memo", str(write(tmp, "memo.md", memo))]
    args += extra or []
    proc = subprocess.run(args, capture_output=True, text=True, encoding="utf-8")
    rep = json.loads((tmp / "g.json").read_text(encoding="utf-8"))
    return proc.returncode, rep, {c["check"]: c["status"] for c in rep["checks"]}


def row(sentence, cue="HEDGE", loc="SRC05 p.12", why="Readers would otherwise extend the finding to tenants without deeds."):
    return {"sentence_excerpt": sentence, "cue_type": cue, "evidence_locator": loc, "why_reader_needs_it": why, "section": "x"}


class ScanTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_zh_traces_and_unanchored_negation(self):
        r = trace_scan.scan(write(self.tmp, "d.md", ZH_DEFENSIVE))
        self.assertEqual(r["lang"], "zh")
        self.assertGreaterEqual(r["totals"]["process_trace"], 2)
        self.assertIn("hedged_finding", {c["type"] for c in r["candidates"]})

    def test_zh_living_material_and_stance(self):
        t = trace_scan.scan(write(self.tmp, "b.md", ZH_BASE))["totals"]
        self.assertGreater(t["quotes_per_k"], 0)
        self.assertGreater(t["numbers_per_k"], 0)
        self.assertGreaterEqual(t["stance"], 2)
        self.assertEqual(t["process_trace"], 0)

    def test_en_scan(self):
        t = trace_scan.scan(write(self.tmp, "e.md", EN_TEXT))["totals"]
        self.assertGreaterEqual(t["process_trace"], 1)
        self.assertGreaterEqual(t["stance"], 1)
        self.assertGreater(t["quotes_per_k"], 0)

    def test_R04_multisentence_quote_is_speaker_voice(self):
        r = trace_scan.scan(write(self.tmp, "q.md", "# 发现\n\n她说：“我们一直等着批复。后来可能要迁走。”她最终留了下来。\n"))
        self.assertFalse([s for s in r["sentences"] if "hedge" in s["flags"]])

    def test_R08_numbered_references_excluded(self):
        body = "# Findings\n\nResidents negotiated access.\n\n# 6 References\n\nSmith, 2020. As recorded in the transcript, households may move.\n"
        self.assertEqual(trace_scan.scan(write(self.tmp, "n.md", body))["totals"]["process_trace"], 0)

    def test_R09_opponent_in_same_paragraph_anchors(self):
        r = trace_scan.scan(write(self.tmp, "a.md", "# Findings\n\nSmith (2020) assumes that allocation follows purchasing power. "
                                                   "The council prioritized residence rather than income.\n"))
        negs = [c for c in r["candidates"] if c["type"] == "negation_defense"]
        self.assertTrue(negs and negs[0]["anchored"])

    def test_cohesion_signals_discriminate(self):
        linked = ("# 3 发现\n\n村里的地块分散在不同坡面上。分散的地块使每户都要单独修水池，"
                  "而单独修水池的成本又让多数农户放弃了滴灌。放弃滴灌的农户于是把精力转到采后加工。\n")
        staccato = ("# 3 发现\n\n村里地块分散。水池很贵。企业有钱。精品靠加工。规则改变一切。\n")
        a = trace_scan.scan(write(self.tmp, "l.md", linked))["totals"]
        b = trace_scan.scan(write(self.tmp, "s.md", staccato))["totals"]
        self.assertGreater(a["topic_link_rate"], b["topic_link_rate"])
        self.assertGreater(b["staccato_runs"], a["staccato_runs"])

    def test_connectives_are_not_closers(self):
        text = ("# 3 发现\n\n企业整片流转了800亩坡地。设施因此按连片地块设计。农户的五亩地无法接入同一管网，"
                "因此他们把投入转向了采后加工。\n")
        r = trace_scan.scan(write(self.tmp, "c.md", text))
        self.assertFalse([c for c in r["candidates"] if c["type"] == "abstract_closer"])


class GateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_defensive_revision_fails(self):
        code, _, st = run_gate(self.tmp, ZH_BASE, ZH_DEFENSIVE)
        self.assertEqual(code, 2)
        self.assertEqual(st["G1_TRACE"], "FAIL")
        self.assertEqual(st["G2_NEW_CUES"], "FAIL")

    def test_rehumanized_revision_passes(self):
        code, rep, st = run_gate(self.tmp, ZH_DEFENSIVE, ZH_BASE)
        self.assertEqual(code, 0, [c for c in rep["checks"] if c["status"] == "FAIL"])

    def test_unchanged_lean_paper_passes(self):
        code, rep, _ = run_gate(self.tmp, ZH_BASE, ZH_BASE)
        self.assertEqual(code, 0, [c for c in rep["checks"] if c["status"] == "FAIL"])

    def test_registered_caution_ok_generic_reason_rejected(self):
        rev = ZH_BASE.replace("本文发现，居民通过延长谈判改变了补偿结果。",
                              "本文发现，居民通过延长谈判改变了补偿结果；这一结果可能只适用于有产权证的住户。")
        _, _, st = run_gate(self.tmp, ZH_BASE, rev, [row("可能只适用于有产权证的住户")])
        self.assertEqual(st["G2_NEW_CUES"], "PASS")
        code, _, st = run_gate(self.tmp, ZH_BASE, rev, [row("可能只适用于有产权证的住户", why="谨慎起见，避免过度推论以防审稿人质疑")])
        self.assertEqual(st["G0_REGISTER"], "FAIL")
        self.assertEqual(st["G2_NEW_CUES"], "FAIL")
        self.assertEqual(code, 2)

    def test_R01_new_hedge_inside_carried_sentence_is_caught(self):
        b = "# Findings\n\nWe argue that the council changed land allocations by revising the eligibility criteria after negotiations with village leaders.\n"
        r = b.replace("changed", "may change")
        code, _, st = run_gate(self.tmp, b, r)
        self.assertEqual(st["G2_NEW_CUES"], "FAIL")
        self.assertEqual(code, 2)

    def test_R02_registered_negation_not_failed_downstream(self):
        sent = "The register measures approval rather than completed delivery."
        code, rep, st = run_gate(self.tmp, "# Findings\n\nThe agency issued written approvals to eligible applicants.\n",
                                 "# Findings\n\nThe agency issued written approvals to eligible applicants. " + sent + "\n",
                                 [row(sent, "NEGATION", why="Without it readers would confuse approval with completed delivery.")])
        self.assertEqual(st["G2_NEW_CUES"], "PASS")
        self.assertEqual(code, 0, [c for c in rep["checks"] if c["status"] == "FAIL"])

    def test_R03_R11_registered_association_in_conclusion_passes(self):
        sent = "The estimate is associated with income."
        why = "The model lacks causal identification, so readers must read the coefficient as an association."
        for head in ("# Abstract\n\n", "# 6 Conclusion\n\n## 6.1 Main finding\n\n"):
            code, rep, st = run_gate(self.tmp, head + "We report the estimate for household income.\n",
                                     head + "We report the estimate for household income. " + sent + "\n",
                                     [row(sent, "HEDGED_FINDING", why=why)])
            self.assertEqual(code, 0, [c for c in rep["checks"] if c["status"] == "FAIL"])
        code, _, st = run_gate(self.tmp, "# 6 Conclusion\n\n## 6.1 Main finding\n\nWe report the estimate for household income.\n",
                               "# 6 Conclusion\n\n## 6.1 Main finding\n\nWe report the estimate for household income. " + sent + "\n")
        self.assertEqual(st["G2_NEW_CUES"], "FAIL")

    def test_R05_trace_after_many_candidates_is_caught(self):
        common = "# Findings\n\n" + "\n\n".join("The interviews show that residents negotiated access." for _ in range(450)) + "\n"
        code, _, st = run_gate(self.tmp, common, common + "\nAs recorded in the transcript, the resident refused the offer.\n")
        self.assertEqual(st["G1_TRACE"], "FAIL")
        self.assertEqual(code, 2)

    def test_R06_empty_revision_fails(self):
        code, _, st = run_gate(self.tmp, "# Findings\n\nWe argue that residents negotiated access to housing.\n", "")
        self.assertEqual(st["G0_INPUT"], "FAIL")
        self.assertEqual(code, 2)

    def test_R07_register_needs_type_and_locator(self):
        sent = "We may have missed households who left before fieldwork."
        code, _, st = run_gate(self.tmp, "# Findings\n\nWe interviewed residents at their homes.\n",
                               "# Findings\n\nWe interviewed residents at their homes. " + sent + "\n",
                               [row(sent, cue="", loc="x", why="Residents who moved away were not in the sampling frame.")])
        self.assertEqual(st["G0_REGISTER"], "FAIL")
        self.assertEqual(code, 2)

    def test_R13_docx_and_md_equivalent(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")
        md = write(self.tmp, "m.md", "# Findings\n\nWe argue that residents negotiated access.\n\n| Item | Number |\n| --- | --- |\n| Sample | The sample contains 50 people. |\n")
        doc = Document()
        doc.add_heading("Findings", level=1)
        doc.add_paragraph("We argue that residents negotiated access.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(1, 1).text = "The sample contains 50 people."
        dp = self.tmp / "m.docx"
        doc.save(dp)
        self.assertEqual(trace_scan.scan(md)["totals"]["numbers_per_k"], trace_scan.scan(dp)["totals"]["numbers_per_k"])
        code, rep, _ = run_gate(self.tmp, dp, md)
        self.assertEqual(code, 0, [c for c in rep["checks"] if c["status"] == "FAIL"])

    def test_quote_loss_needs_memo(self):
        rev = ZH_BASE.replace("“我们家三代都住在这条巷子里，钱再多也买不回邻居。”（R3，2021）", "")
        _, _, st = run_gate(self.tmp, ZH_BASE, rev)
        self.assertEqual(st["G6_LOSS"], "WARN")
        code, _, st = run_gate(self.tmp, ZH_BASE, rev, memo="无关说明")
        self.assertEqual(st["G6_LOSS"], "FAIL")
        _, _, st = run_gate(self.tmp, ZH_BASE, rev, memo="删去“我们家三代都住在这条巷子里”一句：同一论点由第四轮谈判的数字承担。")
        self.assertEqual(st["G6_LOSS"], "PASS")

    def test_cohesion_warning(self):
        linked = ("# 3 发现\n\n村里的地块分散在不同坡面上。分散的地块使每户都要单独修水池，"
                  "而单独修水池的成本又让多数农户放弃了滴灌。放弃滴灌的农户于是把精力转到采后加工。\n")
        staccato = "# 3 发现\n\n村里地块分散在坡面。水池造价很高。企业资金充足。精品依靠加工。规则改变一切。\n"
        _, _, st = run_gate(self.tmp, linked, staccato)
        self.assertEqual(st.get("W8_COHESION"), "WARN")


class ToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_R10_overlap_missing_sibling_is_incomplete(self):
        ms = write(self.tmp, "ms.md", ZH_BASE)
        out = self.tmp / "o.md"
        cmd = [sys.executable, str(SCRIPTS / "sibling_overlap_check.py"), "--manuscript", str(ms), "--out-md", str(out)]
        self.assertEqual(subprocess.run(cmd + ["--sibling", str(self.tmp / "missing.md")], capture_output=True).returncode, 1)
        sib = write(self.tmp, "sib.md", "# 4\n\n她说：“钱再多也买不回邻居，我们三代都住在这里。”（R3，2021）\n")
        self.assertEqual(subprocess.run(cmd + ["--sibling", str(sib)], capture_output=True).returncode, 0)
        sib2 = write(self.tmp, "sib2.md", "引文：“我们家三代都住在这条巷子里，钱再多也买不回邻居。”（R3，2021）")
        self.assertEqual(subprocess.run(cmd + ["--sibling", str(sib2)], capture_output=True).returncode, 2)

    def test_R12_timeline_prefers_dates_in_names(self):
        d = self.tmp / "project"
        old = write(d, "draft_20200101.md", "# Introduction\n\nWe argue that authority is distributed across levels.\n")
        new = write(d, "draft_20260101.md", "# Introduction\n\nWe argue that authority depends on implementation.\n")
        os.utime(old, (2000000000, 2000000000))
        os.utime(new, (1900000000, 1900000000))
        subprocess.run([sys.executable, str(SCRIPTS / "target_drift_timeline.py"), "--root", str(d),
                        "--out-md", str(self.tmp / "t.md"), "--out-json", str(self.tmp / "t.json")], check=True, capture_output=True)
        order = [Path(x["path"]).name for x in json.loads((self.tmp / "t.json").read_text(encoding="utf-8"))]
        self.assertEqual(order[0], "draft_20200101.md")

    def test_final_check_binds_hashes(self):
        final = write(self.tmp, "final.md", ZH_BASE)
        run_gate(self.tmp, ZH_DEFENSIVE, final)
        out = self.tmp / "f.md"
        cmd = [sys.executable, str(SCRIPTS / "final_check.py"), "--manuscript", str(final), "--gate-json", str(self.tmp / "g.json"),
               "--out-md", str(out)]
        self.assertEqual(subprocess.run(cmd, capture_output=True).returncode, 0)
        final.write_text(ZH_BASE + "\n补一句。\n", encoding="utf-8")  # changed after the gate ran
        self.assertEqual(subprocess.run(cmd, capture_output=True).returncode, 2)

    def test_final_check_write_result(self):
        final = write(self.tmp, "final.md", ZH_BASE)
        run_gate(self.tmp, ZH_DEFENSIVE, final)
        result = write(self.tmp, "rrd_result.json", json.dumps({"status": "ROUTE_CORRECTED_READY_FOR_FINAL_POLISH",
                                                                 "final_hashes": {}, "sources_unchanged": None}))
        src = write(self.tmp, "src.txt", "source")
        import hashlib
        manifest = write(self.tmp, "m.csv", "path,sha256\n" + f"{src},{hashlib.sha256(src.read_bytes()).hexdigest()}\n")
        cmd = [sys.executable, str(SCRIPTS / "final_check.py"), "--manuscript", str(final), "--gate-json", str(self.tmp / "g.json"),
               "--result-json", str(result), "--source-manifest", str(manifest), "--out-md", str(self.tmp / "f.md")]
        self.assertEqual(subprocess.run(cmd, capture_output=True).returncode, 2)  # hashes not yet recorded
        self.assertEqual(subprocess.run(cmd + ["--write-result"], capture_output=True).returncode, 0)
        r = json.loads(result.read_text(encoding="utf-8"))
        self.assertTrue(r["sources_unchanged"])
        self.assertEqual(len(r["final_hashes"]), 1)

    def test_overlap_ignores_reference_titles(self):
        ms = write(self.tmp, "ms.md", EN_TEXT)
        body = "# Findings\n\nWe argue that dispatch rules matter.\n\n" + "Filler sentence about couriers. " * 40
        sib = write(self.tmp, "sib.md", body + "\n\n# References\n\nLee, 2020. \"Effective policy implementation in local states and their many varied consequences\". Journal.\n")
        out = self.tmp / "o.json"
        subprocess.run([sys.executable, str(SCRIPTS / "sibling_overlap_check.py"), "--manuscript", str(ms), "--sibling", str(sib),
                        "--out-md", str(self.tmp / "o.md"), "--out-json", str(out)], capture_output=True)
        self.assertEqual(json.loads(out.read_text(encoding="utf-8"))["siblings"][0]["status"], "CHECKED_NO_MATCH")

    def test_rule_debt_skips_interview_speech(self):
        goal = write(self.tmp, "GOAL.md", "不得新增引语。\n受访者说：“需要一个最基础的条件是必须要连片，方便管理”（E1，2025）。\n")
        out = self.tmp / "r.csv"
        subprocess.run([sys.executable, str(SCRIPTS / "rule_debt_extract.py"), str(goal), "--out-csv", str(out)], check=True, capture_output=True)
        with out.open(encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["source_role"], "instruction")


class V3Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_author_paragraph_with_quote_is_not_block(self):
        text = ("# 3 发现\n\n她从小种咖啡，家里二十多亩地。按质论价之后，她说：“我们现在卖一吨抵得上过去三吨。”"
                "她把精力转到采后。“我们自己守着发酵，我们自己看天。”（V2，2025）\n")
        r = trace_scan.scan(write(self.tmp, "b.md", text))
        self.assertEqual(r["sections"]["3 发现"]["block_quotes"] if "block_quotes" in r["sections"]["3 发现"] else 0, 0)
        self.assertEqual(r["totals"]["quotes"], 2)

    def test_relation_wording_is_not_hedged_finding(self):
        r = trace_scan.scan(write(self.tmp, "h.md", "# 3 发现\n\n收购者怎样与种植者相联系，决定了要求能否进入田间。\n"))
        self.assertEqual(r["totals"]["hedged_finding"], 0)

    def test_memo_with_abbreviated_quote_counts(self):
        rev = ZH_BASE.replace("“我们家三代都住在这条巷子里，钱再多也买不回邻居。”（R3，2021）", "一位老住户说，三代人都住在这里。（R3，2021）")
        _, _, st = run_gate(self.tmp, ZH_BASE, rev, memo="改为转述：“我们家三代都住在……买不回邻居”一句，理由是同段已有数字。")
        self.assertEqual(st["G6_LOSS"], "PASS")

    def test_architecture_signals(self):
        cold = ("# 2 研究区域与方法\n\n研究材料来自三轮访谈（R1，2021）。我们访谈了十二人。\n\n"
                "# 3 发现\n\n## 3.1 补偿\n\n补偿机制决定了居民的选择逻辑与组织方式。居民在谈判中改变了结果（R3，2021）。\n\n"
                "补偿安排也影响了社区关系与组织过程。街道干部认为这种做法拖慢了进度（G1，2021）。\n")
        guided = ("# 2 研究区域与方法\n\n本文选取老城区作为案例区，因为这里的补偿争议最集中（R1，2021）。我们访谈了十二人。\n\n"
                  "# 3 发现\n\n## 3.1 补偿\n\n2019年春天，一位住了三代的老住户在第四轮谈判后才签字（R3，2021）。她的坚持改变了补偿结果。\n\n"
                  "与这位老住户不同，街道干部看到的是工期，他认为这种做法拖慢了进度（G1，2021）。\n")
        a = trace_scan.scan(write(self.tmp, "c.md", cold))
        b = trace_scan.scan(write(self.tmp, "g.md", guided))
        self.assertGreater(a["totals"]["abstract_openers"], b["totals"]["abstract_openers"])
        self.assertEqual(a["totals"]["unguided_shifts"], 1)
        self.assertEqual(b["totals"]["unguided_shifts"], 0)
        self.assertFalse(a["case_selection_sections"])
        self.assertTrue(b["case_selection_sections"])
        _, rep, st = run_gate(self.tmp, write(self.tmp, "gb.md", guided), write(self.tmp, "cr.md", cold))
        self.assertEqual(st.get("W13_SECTION_ARCHITECTURE"), "WARN")


class V31Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_short_attributed_quote_counts(self):
        r = trace_scan.scan(write(self.tmp, "s.md", "# 3 发现\n\n讲到缺水时，他提醒：“你看表层没有用”（R2，2025）。\n"))
        self.assertEqual(r["totals"]["quotes"], 1)

    def test_markdown_title_is_not_a_chapter(self):
        text = ("# 论文题名\n\n摘要：一句话。\n\n## 2 研究区域与方法\n\n### 2.1 案例\n\n本文选取老城区作为案例区，"
                "因为争议最集中（R1，2021）。\n\n## 3 发现\n\n### 3.1 补偿\n\n2019年，一位老住户签了字（R3，2021）。\n")
        r = trace_scan.scan(write(self.tmp, "t.md", text))
        self.assertIn("2 研究区域与方法", r["chapter_order"])
        self.assertEqual(r["sections"]["2.1 案例"]["chapter"], "2 研究区域与方法")

    def test_material_renewal_warning(self):
        _, _, st = run_gate(self.tmp, ZH_BASE, ZH_BASE)
        self.assertEqual(st.get("W14_MATERIAL_RENEWAL"), "WARN")
        rev = ZH_BASE.replace("本文发现，居民通过延长谈判改变了补偿结果。",
                              "一位街道干部说：“拖到第四轮，补偿就涨了。”（G1，2021）本文发现，居民通过延长谈判改变了补偿结果。")
        _, _, st = run_gate(self.tmp, ZH_BASE, rev)
        self.assertNotIn("W14_MATERIAL_RENEWAL", st)

    def test_semantic_review_must_name_siblings(self):
        final = write(self.tmp, "final.md", ZH_BASE)
        run_gate(self.tmp, ZH_DEFENSIVE, final)
        ov = write(self.tmp, "ov.json", json.dumps({"manuscript_sha256": __import__("hashlib").sha256(final.read_bytes()).hexdigest(),
                                                    "overall": "SEMANTIC_REVIEW_REQUIRED",
                                                    "siblings": [{"file": "C:/x/english_sibling_paper.md", "status": "SEMANTIC_REVIEW"}]}))
        cmd = [sys.executable, str(SCRIPTS / "final_check.py"), "--manuscript", str(final), "--gate-json", str(self.tmp / "g.json"),
               "--overlap-json", str(ov), "--out-md", str(self.tmp / "f.md")]
        bad = write(self.tmp, "rev_bad.md", "reviewed everything, all fine " * 20)
        self.assertEqual(subprocess.run(cmd + ["--semantic-review", str(bad)], capture_output=True).returncode, 2)
        good = write(self.tmp, "rev_good.md", "english_sibling_paper: 12 quotations compared item by item; none reused.")
        self.assertEqual(subprocess.run(cmd + ["--semantic-review", str(good)], capture_output=True).returncode, 0)


class V32Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_conversion_template_warning(self):
        plain = ("# 3 发现\n\n2019年，一位老住户在第四轮谈判后签字（R3，2021）。她坚持了三个月，补偿涨了三成。"
                 "街道干部说这拖慢了工期（G1，2021）。\n")
        templ = ("# 3 发现\n\n2019年，一位老住户在第四轮谈判后签字（R3，2021）。谈判使补偿成为可协商的对象。"
                 "坚持让时间进入定价。街道干部说这拖慢了工期（G1，2021）。工期由此转为谈判筹码。\n")
        _, _, st = run_gate(self.tmp, plain, templ)
        self.assertEqual(st.get("W15_CONVERSION_TEMPLATE"), "WARN")

    def test_prose_paragraph_starting_with_figure_reference_is_kept(self):
        text = ("# 1 框架\n\n图1据此连接三个分析问题：什么被选择，怎样被使用，由谁承担。下文先追踪供水，再转入加工。最后比较合作。\n\n"
                "图1 分析框架\n")
        r = trace_scan.scan(write(self.tmp, "f.md", text))
        self.assertEqual(r["totals"]["paragraphs"], 1)

    def test_timeline_file_list_ignores_default_excludes(self):
        d = self.tmp / "proj" / "source_texts"
        f = write(d, "draft_20240101.md", "# Introduction\n\nWe argue that land is assembled.\n")
        lst = write(self.tmp, "list.txt", str(f))
        subprocess.run([sys.executable, str(SCRIPTS / "target_drift_timeline.py"), "--root", str(self.tmp / "proj"),
                        "--file-list", str(lst), "--out-md", str(self.tmp / "t.md"), "--out-json", str(self.tmp / "t.json")],
                       check=True, capture_output=True)
        self.assertEqual(len(json.loads((self.tmp / "t.json").read_text(encoding="utf-8"))), 1)


class V33Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_relation_verb_tic_warning(self):
        plain = ("# 3 发现\n\n2019年，一位老住户在第四轮谈判后签字（R3，2021）。她坚持了三个月，补偿涨了三成。"
                 "街道干部说这拖慢了工期（G1，2021）。\n")
        tic = ("# 3 发现\n\n2019年，一位老住户在第四轮谈判后签字（R3，2021）。补偿要求由此进入谈判。"
               "家庭关系接入定价。街道干部说这拖慢了工期（G1，2021）。工期因此纳入协商，居民参与构成补偿规则。\n")
        r = trace_scan.scan(write(self.tmp, "t.md", tic))
        self.assertGreater(r["totals"]["relation_verbs_per_k"], 1.5)
        self.assertTrue(r["top_relation_verbs"])
        _, _, st = run_gate(self.tmp, plain, tic)
        self.assertEqual(st.get("W16_RELATION_VERB_TIC"), "WARN")
        _, _, st = run_gate(self.tmp, tic, plain)
        self.assertNotIn("W16_RELATION_VERB_TIC", st)

    def test_final_check_requires_explicit_gate_status_and_accepts_two_reports(self):
        final_md = write(self.tmp, "final.md", ZH_BASE)
        run_gate(self.tmp, ZH_DEFENSIVE, final_md)
        g1 = self.tmp / "g1.json"
        g1.write_text((self.tmp / "g.json").read_text(encoding="utf-8"), encoding="utf-8")
        base = [sys.executable, str(SCRIPTS / "final_check.py"), "--manuscript", str(final_md), "--out-md", str(self.tmp / "f.md")]
        self.assertEqual(subprocess.run(base + ["--gate-json", str(g1), "--gate-json", str(self.tmp / "g.json")],
                                        capture_output=True).returncode, 0)
        g = json.loads(g1.read_text(encoding="utf-8"))
        g.pop("overall", None)
        g1.write_text(json.dumps(g), encoding="utf-8")
        self.assertEqual(subprocess.run(base + ["--gate-json", str(g1)], capture_output=True).returncode, 2)

    def test_final_check_write_result_records_status(self):
        final = write(self.tmp, "final.md", ZH_BASE)
        run_gate(self.tmp, ZH_DEFENSIVE, final)
        result = write(self.tmp, "rrd_result.json", json.dumps({"status": "ROUTE_CORRECTED_READY_FOR_FINAL_POLISH", "final_hashes": {}}))
        subprocess.run([sys.executable, str(SCRIPTS / "final_check.py"), "--manuscript", str(final), "--gate-json", str(self.tmp / "g.json"),
                        "--result-json", str(result), "--write-result", "--out-md", str(self.tmp / "f.md")], capture_output=True)
        self.assertEqual(json.loads(result.read_text(encoding="utf-8"))["final_check"], "PASS")

    def test_overlap_checks_short_quotes_exactly(self):
        ms = write(self.tmp, "ms.md", "# 3 发现\n\n讲到缺水时，他提醒：“你看表层没有用”（R2，2025）。\n")
        sib = write(self.tmp, "sib.md", "# 发现\n\n他说：“你看表层没有用”（R2，2025），随后打开了记录仪。\n")
        out = self.tmp / "o.json"
        subprocess.run([sys.executable, str(SCRIPTS / "sibling_overlap_check.py"), "--manuscript", str(ms), "--sibling", str(sib),
                        "--out-md", str(self.tmp / "o.md"), "--out-json", str(out)], capture_output=True)
        self.assertEqual(json.loads(out.read_text(encoding="utf-8"))["overall"], "MATCHES")
        self.assertIn("你看表层没有用", (self.tmp / "o.md").read_text(encoding="utf-8"))

    def test_timeline_reports_whitelisted_files_not_read(self):
        f = write(self.tmp / "proj", "draft_20240101.md", "# Introduction\n\nWe argue that land is assembled.\n" * 50)
        lst = write(self.tmp, "list.txt", f"{f}\n{self.tmp / 'proj' / 'missing.md'}\n")
        subprocess.run([sys.executable, str(SCRIPTS / "target_drift_timeline.py"), "--root", str(self.tmp / "proj"),
                        "--file-list", str(lst), "--max-bytes", "100", "--out-md", str(self.tmp / "t.md")],
                       check=True, capture_output=True)
        md = (self.tmp / "t.md").read_text(encoding="utf-8")
        self.assertIn("Files not read", md)
        self.assertIn("max-bytes", md)
        self.assertIn("missing", md)


class V34Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_tag_word_surge(self):
        para = "2019年，一位老住户在第四轮谈判后签字（R{n}，2021）。她坚持了三个月，补偿涨了三成。\n\n"
        plain = "# 3 发现\n\n" + "".join(para.format(n=i) for i in range(12))
        tagged = "# 3 发现\n\n" + "".join(
            para.format(n=i).replace("补偿涨了三成。", "补偿涨了三成，谈判由此获得新的筹码，筹码再获得用途。") for i in range(12))
        def tagged_words(rep):
            c = next((c for c in rep["checks"] if c["check"] == "W17_TAG_WORD"), None)
            return {i.split(":")[0] for i in c["items"]} if c else set()
        _, rep, st = run_gate(self.tmp, plain, tagged)
        self.assertIn("用途", tagged_words(rep))
        _, _, st = run_gate(self.tmp, plain, plain)
        self.assertNotIn("W17_TAG_WORD", st)
        _, rep, _ = run_gate(self.tmp, plain, tagged, extra=["--core-term", "用途"])
        self.assertNotIn("用途", tagged_words(rep))

    def test_internal_document_labels_are_traces(self):
        r = trace_scan.scan(write(self.tmp, "l.md", "# 2 方法\n\n本研究与论文①共享田野，姊妹论文的结论不进入本稿。\n"))
        self.assertGreaterEqual(r["totals"]["process_trace"], 1)

    def test_case_selection_with_reason_clause(self):
        r = trace_scan.scan(write(self.tmp, "c.md", "# 2 研究区域与方法\n\n本文选择老城区，是因为这里的补偿争议最集中（R1，2021）。\n"))
        self.assertTrue(r["case_selection_sections"])

    def test_condition_clause_guides_focus_shift(self):
        text = ("# 3 发现\n\n## 3.1 补偿\n\n2019年，一位老住户在第四轮谈判后签字（R3，2021）。她坚持了三个月。\n\n"
                "地块分散时，街道干部要跑更多的门，他认为这拖慢了进度（G1，2021）。\n\n"
                "除了工期，资金也是问题，一位会计说账上只剩两成（F1，2021）。\n")
        r = trace_scan.scan(write(self.tmp, "m.md", text))
        self.assertEqual(r["totals"]["unguided_shifts"], 0)


if __name__ == "__main__":
    unittest.main()
