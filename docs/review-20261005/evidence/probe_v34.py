#!/usr/bin/env python3
"""Targeted RRD v3.4 validation probes. Synthetic fixtures; no model invocation.
Run: python probe_v34.py --skill-root PATH --out DIRECTORY
The expected outcomes encode proposed fail-closed acceptance requirements,
not assertions that the package already implements each requirement.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, os, subprocess, sys
from pathlib import Path

def main() -> int:
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--skill-root',type=Path,required=True)
    ap.add_argument('--out',type=Path,required=True)
    a=ap.parse_args(); root=a.skill_root.resolve(); out=a.out.resolve(); out.mkdir(parents=True,exist_ok=True)
    env=dict(os.environ,PYTHONDONTWRITEBYTECODE='1'); results=[]
    def write(p:Path,text:str)->Path:
        p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding='utf8');return p
    def js(p:Path,data)->Path: return write(p,json.dumps(data,ensure_ascii=False,indent=2))
    def sha(p:Path)->str:return hashlib.sha256(p.read_bytes()).hexdigest()
    def invoke(cid,script,opts,expect,requirement):
        d=out/cid;d.mkdir(exist_ok=True);report=d/'report.json'
        cmd=[sys.executable,str(root/'scripts'/script),*map(str,opts),'--out-json',str(report),'--out-md',str(d/'report.md')]
        proc=subprocess.run(cmd,capture_output=True,text=True,encoding='utf8',env=env,timeout=20)
        write(d/'stdout.txt',proc.stdout);write(d/'stderr.txt',proc.stderr)
        r=json.loads(report.read_text()) if report.exists() else {}
        actual=r.get('overall','NO_REPORT');ok=(actual=='FAIL') if expect=='FAIL' else (actual in {'PASS','PASS_WITH_WARNINGS'})
        rec={'id':cid,'requirement':requirement,'expected':expect,'actual':actual,'exit_code':proc.returncode,'meets_requirement':ok,'report':str(report.relative_to(out)),'command':cmd}
        results.append(rec); print(cid,actual,'MET' if ok else 'GAP',flush=True)
        return r
    b=write(out/'source.md','# Findings\n\nThe council revised eligibility after residents submitted their records.\n')
    md=write(out/'final.md',b.read_text())
    # A second format with identical prose; formatting is immaterial to this contract test.
    from docx import Document
    doc=Document();doc.add_heading('Findings',level=1);doc.add_paragraph('The council revised eligibility after residents submitted their records.');docp=out/'final.docx';doc.save(docp)
    gm=js(out/'gate_md.json',{'overall':'PASS','revised':{'sha256':sha(md)}})
    gd=js(out/'gate_docx.json',{'overall':'PASS','revised':{'sha256':sha(docp)}})
    # Stub gate reports isolate the final_check interface; they are not evidence of manuscript quality.
    common=['--manuscript',md,'--gate-json',gm]
    invoke('A01_ungated_final','final_check.py',common+['--manuscript',docp],'FAIL','Every supplied final artifact must be covered by a gate report, not merely every report matched to some artifact.')
    invoke('C01_complete_coverage','final_check.py',common+['--manuscript',docp,'--gate-json',gd],'PASS','Both final artifacts have matching gate reports.')
    empty=write(out/'empty_manifest.csv','path,sha256\n')
    valid=write(out/'valid_manifest.csv',f'path,sha256\n{b},{sha(b)}\n')
    invoke('A02_empty_manifest','final_check.py',common+['--source-manifest',empty],'FAIL','An empty input manifest must not certify that all workflow sources were verified unchanged.')
    invoke('C02_populated_manifest','final_check.py',common+['--source-manifest',valid],'PASS','A populated unchanged source manifest is accepted.')
    bogus=js(out/'bogus_result.json',{'status':'UNRECOGNIZED_STATE','final_hashes':{'final.md':sha(md)}})
    ready=js(out/'ready_result.json',{'status':'ROUTE_CORRECTED_READY_FOR_FINAL_POLISH','final_hashes':{'final.md':sha(md)}})
    invoke('A03_unknown_status','final_check.py',common+['--result-json',bogus],'FAIL','Result status should be validated against the documented enum; a nonempty arbitrary string is not a valid workflow state.')
    invoke('C03_recognized_status','final_check.py',common+['--result-json',ready],'PASS','A documented completion status is accepted.')
    rev=write(out/'multicue.md','# Findings\n\nThe council may have revised eligibility rather than budgets after residents submitted their records.\n')
    sentence='The council may have revised eligibility rather than budgets after residents submitted their records.'
    def register(p:Path,kinds):
        with p.open('w',encoding='utf8',newline='') as f:
            w=csv.DictWriter(f,fieldnames=['sentence_excerpt','cue_type','evidence_locator','why_reader_needs_it','section']);w.writeheader()
            for k in kinds:w.writerow({'sentence_excerpt':sentence,'cue_type':k,'evidence_locator':'SYNTHETIC-SRC01 paragraph 1','why_reader_needs_it':'The available record leaves the timing of the eligibility revision uncertain.' if k=='HEDGE' else 'The recorded decision concerns eligibility rules, not a change in budget allocations.','section':'Findings'})
        return p
    one=register(out/'one_type.csv',['HEDGE']);two=register(out/'two_types.csv',['HEDGE','NEGATION'])
    gateargs=['--baseline',b,'--revised',rev]
    invoke('A04_partial_cue_register','defense_gate.py',gateargs+['--register',one],'FAIL','Each newly added cue type requires its own matching registration; HEDGE must not authorize an unregistered NEGATION.')
    invoke('C04_complete_cue_register','defense_gate.py',gateargs+['--register',two],'PASS','Each new cue type is registered separately.')
    invoke('A05_missing_explicit_config','defense_gate.py',['--baseline',b,'--revised',md,'--config',out/'absent_config.json'],'FAIL','An explicitly requested but missing configuration must not silently fall back to defaults.')
    invoke('C05_existing_config','defense_gate.py',['--baseline',b,'--revised',md,'--config',root/'assets/templates/gate_config.default.json'],'PASS','An existing configuration is accepted.')
    template=json.loads((root/'assets/templates/rrd_result.template.json').read_text());template['status']='ROUTE_CORRECTED_READY_FOR_FINAL_POLISH'
    rp=js(out/'result_from_template.json',template)
    r=invoke('C06_write_result','final_check.py',common+['--result-json',rp,'--write-result','--source-manifest',valid],'PASS','Documented write-result execution completes; schema compatibility is examined separately.')
    after=json.loads(rp.read_text())
    schema={'id':'A06_result_schema','requirement':'final_check field preserves the object structure declared by rrd_result.template.json','before_type':type(template['final_check']).__name__,'after_type':type(after['final_check']).__name__,'after_value':after['final_check'],'meets_requirement':type(template['final_check']) is type(after['final_check'])}
    summary={'scope':'synthetic contract probes; not a manuscript-quality evaluation','results':results,'schema_observation':schema,'total_executed_cases':len(results),'control_cases':sum(x['id'].startswith('C') for x in results),'gap_cases':sum(not x['meets_requirement'] for x in results)}
    js(out/'summary.json',summary)
    print(json.dumps(schema,ensure_ascii=False),flush=True)
    return 0
if __name__=='__main__':raise SystemExit(main())
